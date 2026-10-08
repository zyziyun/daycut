"""vstudio.intake: golden prompts x material mixes through the rule planner and a mocked model (validation,
defaults, focus ranges, checkpoints, estimates), revise, apply (projects + series), the inventory on synthetic
media / documents, and the CLI. No network, no ASR, no mediapipe (hooks replaced)."""
import json
import os
import shutil
import subprocess
import sys

import pytest

from vstudio.intake import apply as AP
from vstudio.intake import docs as D
from vstudio.intake import inventory as I
from vstudio.intake import plan as PL
from vstudio.intake import rules as R
from vstudio.project.core import Project

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
HAS_FFMPEG = bool(shutil.which("ffmpeg"))


@pytest.fixture(autouse=True)
def _iso(tmp_path, monkeypatch):
    monkeypatch.setenv("VSTUDIO_HOME", str(tmp_path / "home"))
    monkeypatch.setenv("VSTUDIO_CACHE", str(tmp_path / "cache"))
    monkeypatch.setenv("VSTUDIO_BATCH_BENCH", str(tmp_path / "bench.json"))
    for k in ("ANTHROPIC_API_KEY", "OPENAI_API_KEY", "VSTUDIO_LLM_PROVIDER", "VSTUDIO_LLM_INTAKE_PROVIDER",
              "VSTUDIO_CLIENTS"):
        monkeypatch.delenv(k, raising=False)


# --------------------------------------------------------------------------- fake analyses (no media work)
class Mat:
    """Builds an analysis dict like inventory.analyze would, with real (empty) files so apply can create projects."""

    def __init__(self, d, template=None):
        self.d, self.files, self.groups, self.template = d, [], [], template

    def _add(self, name, kind, **facts):
        p = self.d / name
        p.parent.mkdir(parents=True, exist_ok=True)
        if kind == "video" and self.template:
            shutil.copy(self.template, p)                  # a real (tiny) video: project planning probes it
        else:
            p.write_bytes(b"x")
        f = dict(id=f"f{len(self.files) + 1}", path=str(p), rel=name, kind=kind, size=1, qhash=f"h{len(self.files)}",
                 **facts)
        self.files.append(f)
        return f

    def video(self, name, dur, orient="horizontal", **kw):
        w, h = (1920, 1080) if orient == "horizontal" else (1080, 1920)
        base = dict(duration=dur, width=w, height=h, orientation=orient, has_audio=True, speech=True, language="zh",
                    talking_head=False, multi_person=False, screen_share=False, burned_captions=False)
        base.update(kw)
        return self._add(name, "video", **base)

    def photo(self, name, **kw):
        return self._add(name, "image", width=4000, height=3000, orientation="horizontal", **kw)

    def audio(self, name, dur, speech=False):
        return self._add(name, "audio", duration=dur, has_audio=True, speech=speech)

    def doc(self, name, shape="article", headings=(), title=None, **kw):
        return self._add(name, "text", chars=5000, lang="zh", shape=shape, headings=list(headings),
                         title=title or os.path.splitext(name)[0], excerpt="……", urls=[], **kw)

    def transcript(self, f, sentences):
        tp = self.d / f"{f['id']}.transcript.json"
        tp.write_text(json.dumps(dict(sentences=sentences), ensure_ascii=False), encoding="utf-8")
        f["transcript"] = str(tp)

    def group(self, kind, ids, **kw):
        self.groups.append(dict(id=f"g{len(self.groups) + 1}", kind=kind, folder=str(self.d), files=ids, n=len(ids), **kw))

    def analysis(self):
        tot = dict(videos=sum(f["kind"] == "video" for f in self.files),
                   video_s=sum(f.get("duration", 0) for f in self.files if f["kind"] == "video"),
                   audio=sum(f["kind"] == "audio" for f in self.files), audio_s=0.0,
                   images=sum(f["kind"] == "image" for f in self.files), texts=sum(f["kind"] == "text" for f in self.files),
                   other=0)
        return dict(version=I.VERSION, kind="vstudio.intake.analysis", inputs=[str(self.d)], digest="d" * 16,
                    asr="sample", totals=tot, files=self.files, groups=self.groups, urls=[], notes=[])


def _sents(spans):
    """[(t0, t1, text)] -> transcript sentences"""
    return [dict(t=a, te=b, text=t) for a, b, t in spans]


def _career_transcript(dur=650.0):
    out, t = [], 0.0
    filler = ["今天跟大家复盘一下我的几个副业", "先说说我做过的一些兼职", "这个过程其实挺辛苦的", "后来我换了一个方向"]
    while t < dur - 10:
        text = filler[int(t) % len(filler)]
        if 280 <= t < 360:
            text = "我最大的感悟是副业让我学到了很多新的东西，其实这种成长比收入更重要"
        if 520 <= t < 620:
            text = "做自媒体这件事让我认识了很多博主，我觉得自媒体最重要的是真诚的交流"
        out.append((round(t, 1), round(t + 7.5, 1), text))
        t += 8.0
    return _sents(out)


def build(case, d, template=None):
    m = Mat(d, template)
    if case == "lesson":
        m.video("第三课-上课实录.mp4", 5400, screen_share=True, name_hints=["lecture"])
    elif case == "talk":
        m.video("talk.mov", 185, orient="vertical", talking_head=True)
    elif case == "finished":
        f = m.video("多元副业复盘_final.mp4", 653, talking_head=True, burned_captions=True, name_hints=["final-export"])
        m.transcript(f, _career_transcript())
    elif case == "photos":
        ids = [m.photo(f"trip/IMG_{k:03d}.jpg", taken=f"2026-05-0{1 + k % 3} 10:00:00")["id"] for k in range(12)]
        for k in range(3):
            ids.append(m.video(f"trip/clip{k}.mov", 12, orient="vertical", speech=False)["id"])
        m.group("photo-set", ids[:12], taken=["2026-05-01 10:00:00", "2026-05-03 10:00:00"])
    elif case == "pdf":
        m.doc("transformer.pdf", headings=["什么是注意力", "位置编码", "多头注意力", "残差与归一化", "训练技巧", "推理加速"],
              pages=24)
    elif case == "drama":
        m.doc("剧本.docx", shape="script", episodes=6, title="海边来信")
    elif case == "podcast":
        m.video("zoom_GMT20260930-120000_Recording.mp4", 3600, multi_person=True, faces_median=3, name_hints=["call"])
    elif case == "mixed":
        m.video("course/上课录屏.mp4", 4800, screen_share=True, name_hints=["lecture"])
        m.doc("course/讲义.pdf", headings=["第一章 选题", "第二章 脚本", "第三章 拍摄"], pages=30)
        m.doc("course/notes.md", shape="notes", headings=["开场", "三个误区"])
    elif case == "vlog":
        for k in range(6):
            m.video(f"dji/DJI_{k:04d}.MP4", 40, speech=False)
        m.audio("bgm.mp3", 120)
    elif case == "polish":
        m.video("export_final.mp4", 95, orient="vertical", talking_head=True, burned_captions=True)
    elif case == "course":
        m.video("rec1.mp4", 2400, screen_share=True, name_hints=["screen-recording"])
        m.video("rec2.mp4", 2100, screen_share=True, name_hints=["screen-recording"])
    elif case == "notes":
        m.doc("ideas.md", shape="notes", headings=["AI 求职", "副业", "英语"])
    return m.analysis()


# (case, prompt, expectations)
GOLDEN = [
    ("lesson", "把这节课切成 20 条竖屏", dict(recipes=["longform-to-short"], method="planner",
                                            params={"count": 20, "layout_mode": "split", "platforms": ["xiaohongshu:vertical"]})),
    ("talk", "这段口播剪干净发小红书", dict(recipes=["talkinghead"], method="per-file",
                                          params={"platforms": ["xiaohongshu:vertical"], "cleanup_profile": "standard",
                                                  "speed": 1.25})),     # the talking-head format, not the generic 1.1
    ("finished", "把后面自媒体的思考单独剪出来", dict(recipes=["talkinghead"], method="focus",
                                                      params={"layout": "band", "crop_bottom": 0.28, "captions": True,
                                                              "speed": 1.0, "cleanup_profile": "gentle"},
                                                      ranges_after=450)),
    ("photos", "这些照片和视频做个文艺片", dict(recipes=["photo-story"], method="single", inputs=["photos", "clips"])),
    ("pdf", "用这份 PDF 做 5 条讲解短视频", dict(recipes=["explainer"], method="list", count=5, params={"mode": "short"})),
    ("drama", "用这个剧本做 AI 短剧第一季 6 集", dict(recipes=["ai-video"], method="episodes", count=6,
                                                     needs=["budget", "script"])),
    ("podcast", "播客里有意思的部分剪出来，嘉宾遮脸", dict(recipes=["call-clips"], params={"no_mask": False},
                                                         needs=["consent"], questions=1)),
    ("mixed", "这些素材都帮我做成内容", dict(recipes=["longform-to-short", "explainer", "preproduction"], series=True)),
    ("vlog", "旅行素材剪个卡点 vlog 发抖音", dict(recipes=["vlog"], params={"style": "fun", "platforms": ["douyin"]},
                                                 inputs=["footage", "music"])),
    ("polish", "导出后收尾一下，第一帧黑了，响度也处理下", dict(recipes=["polish"], method="per-file")),
    ("course", "把这两个录屏剪成课程，发 B站", dict(recipes=["longform-course"], params={"platforms": ["bilibili"]})),
    ("notes", "用这些笔记写 3 条口播稿", dict(recipes=["preproduction"], method="list", count=3)),
]


def _check(plan, exp):
    got = [p["recipe"] for p in plan["projects"]]
    assert got == exp["recipes"], (got, plan.get("warnings"), plan.get("risks"))
    p = plan["projects"][0]
    if "method" in exp:
        assert p["items"]["method"] == exp["method"]
    if "count" in exp:
        assert p["items"]["count"] == exp["count"]
    for k, v in (exp.get("params") or {}).items():
        assert p["params"].get(k) == v, (k, p["params"])
    for k in exp.get("inputs") or []:
        assert k in p["inputs"], p["inputs"]
    for cid in exp.get("needs") or []:
        assert any(c["id"] == cid and c["needs_you"] for c in p["checkpoints"]), p["checkpoints"]
    if exp.get("ranges_after"):
        rows = p["items"]["rows"]
        assert rows and all(r["params"]["range"][0] >= exp["ranges_after"] for r in rows), rows
        assert all(r["inputs"].get("video") for r in rows)
    if exp.get("questions"):
        assert len(plan["questions"]) >= exp["questions"]
    if exp.get("series"):
        assert plan["series"] and plan["series"]["id"].startswith("intake-")
    assert PL.validate(plan) == [], PL.validate(plan)
    assert plan["summary_zh"] and plan["estimate"]["machine_min"] >= 0
    for p in plan["projects"]:
        assert p["estimate"]["basis"]["items"] >= 1
        assert p["checkpoints"] and any(c["kind"] == "publish" for c in p["checkpoints"]) or p["recipe"] in (
            "preproduction",)


@pytest.mark.parametrize("case,prompt,exp", GOLDEN, ids=[g[0] for g in GOLDEN])
def test_golden_rule_planner(tmp_path, case, prompt, exp):
    plan = PL.make_plan(prompt, analysis=build(case, tmp_path), asr="off")
    assert plan["planner"]["fallback"] is True and plan["planner"]["provider"] == "none"
    _check(plan, exp)


def _model_reply(case, analysis):
    """What a well-behaved model answers for each golden case (material ids, rows with ranges ...)."""
    f1 = analysis["files"][0]["id"]
    rep = {
        "lesson": [dict(recipe="longform-to-short", name="第三课切片", inputs={"source": [f1]},
                        items=dict(method="planner", count=20), params={"platforms": ["xiaohongshu:vertical"]})],
        "talk": [dict(recipe="talkinghead", inputs={"video": [f1]}, items=dict(method="per-file"),
                      params={"cleanup_profile": "standard"})],
        "finished": [dict(recipe="talkinghead", name="自媒体思考", inputs={"video": [f1]},
                          items=dict(method="focus", rows=[dict(id="zmt1", title="自媒体最重要的是真诚", range=[520, 610]),
                                                           dict(id="bad", range=[900, 990])]),
                          params={"captions": False, "speed": 1.0, "cleanup_profile": "gentle", "bogus": 1})],
        "photos": [dict(recipe="photo-story", materials=["g1"], inputs={"photos": ["g1"], "clips": ["f13", "f14", "f15"]},
                        items=dict(method="single"))],
        "pdf": [dict(recipe="explainer", materials=[f1], items=dict(method="list", rows=[
            dict(id=f"e{k}", inputs={"topic": h}, title=h) for k, h in enumerate((analysis["files"][0].get("headings") or [])[:5])]),
            params={"mode": "short", "platforms": ["xiaohongshu"]})],
        "drama": [dict(recipe="ai-video", materials=[f1], items=dict(method="episodes", count=6))],
        "podcast": [dict(recipe="call-clips", inputs={"source": f1}, items=dict(method="planner", count=5),
                         params={"no_mask": False})],
        "mixed": [dict(recipe="longform-to-short", inputs={"source": "f1"}, items=dict(method="planner", count=12)),
                  dict(recipe="explainer", materials=["f2"], items=dict(method="list", rows=[dict(id="c1", inputs={"topic": "选题"})])),
                  dict(recipe="preproduction", inputs={"script": ["f3"]}, items=dict(method="per-file")),
                  dict(recipe="made-up-recipe", inputs={})],
        "vlog": [dict(recipe="vlog", inputs={"footage": [f["id"] for f in analysis["files"] if f["kind"] == "video"],
                                             "music": "f7"}, params={"style": "fun"}, items=dict(method="single"))],
        "polish": [dict(recipe="polish", inputs={"video": [f1]}, items=dict(method="per-file"))],
        "course": [dict(recipe="longform-course", inputs={"source": ["f1", "f2"]}, items=dict(method="per-file"))],
        "notes": [dict(recipe="preproduction", items=dict(method="list", rows=[
            dict(id=f"n{k}", inputs={"topic": t}) for k, t in enumerate(["AI 求职", "副业", "英语"])]))],
    }[case]
    return dict(projects=rep, questions=[dict(project=0, text="嘉宾都遮吗？", options=["都遮", "不遮"], default="都遮")]
                if case == "podcast" else [], risks=[], summary_zh=f"模型摘要：{case}")


@pytest.mark.parametrize("case,prompt,exp", GOLDEN, ids=[g[0] for g in GOLDEN])
def test_golden_mocked_model(tmp_path, case, prompt, exp):
    analysis = build(case, tmp_path)
    seen = {}

    def call(system, body):
        seen["system"], seen["body"] = system, body
        return dict(json=_model_reply(case, analysis), model="mock", cost_usd=0.0)

    plan = PL.make_plan(prompt, analysis=analysis, asr="off", call=call)
    assert plan["planner"]["fallback"] is False
    assert "never invent a recipe" in seen["system"] and "talkinghead" in seen["body"]
    _check(plan, exp)
    assert plan["summary_zh"] == f"模型摘要：{case}"
    if case == "finished":
        assert "TRANSCRIPT of f1" in seen["body"]
        rows = plan["projects"][0]["items"]["rows"]
        assert [r["id"] for r in rows] == ["zmt1"]                 # 900 s is past the 653 s video: dropped
        assert any("bogus" in w for w in plan["warnings"])
    if case == "mixed":
        assert any("made-up-recipe" in w for w in plan["warnings"])


def test_model_garbage_falls_back(tmp_path):
    a = build("talk", tmp_path)
    plan = PL.make_plan("这段口播剪干净发小红书", analysis=a, asr="off",
                        call=lambda s, b: dict(json=dict(projects=[dict(recipe="nope")]), model="mock"))
    assert plan["planner"]["fallback"] is True and plan["projects"][0]["recipe"] == "talkinghead"

    def boom(s, b):
        raise RuntimeError("401 token expired")
    plan = PL.make_plan("这段口播剪干净发小红书", analysis=a, asr="off", call=boom)
    assert plan["planner"]["fallback"] is True and "401" in plan["planner"]["reason"]


# --------------------------------------------------------------------------- prompt parsing / focus
def test_parse_prompt():
    it = R.parse_prompt("把后面对于自媒体的思考，以及中间的一些感悟剪出来，单独发小红书，每条 60 秒内，1.2倍")
    assert it["extract"] and it["platforms"] == ["xiaohongshu"] and it["max_s"] == 60 and it["speed"] == 1.2
    tops = {f["where"]: f["terms"] for f in it["focus"]}
    assert "自媒体" in tops["后面"][0] and "感悟" in tops["中间"]
    assert R.parse_prompt("用这个剧本做 AI 短剧第一季 6 集")["episodes"] == 6
    assert R.parse_prompt("把这节课切成二十条")["count"] == 20
    assert R.parse_prompt("不要讲解视频")["exclude"] == ["explainer"]
    it = R.parse_prompt("发视频号和小红书，还有快手")
    assert it["platforms"] == ["xiaohongshu", "wechat-channels"] and it["unsupported_platforms"] == ["快手"]   # registry order
    # international first whatever order the request names them in (the creator's rule, everywhere)
    assert R.parse_prompt("发抖音、小红书和 TikTok")["platforms"] == ["tiktok", "xiaohongshu", "douyin"]
    assert sorted(R.parse_prompt("发布方案还要支持 X 和 ins")["platforms"]) == ["instagram", "x"]
    assert R.parse_prompt("1.2x 速度，发 B站")["platforms"] == ["bilibili"]


def test_focus_ranges_respect_position():
    sents = _career_transcript()
    picks = R.focus_ranges(sents, 650, R.focus_of("把后面对于自媒体的思考剪出来"), count=2, max_s=120)
    assert picks and all(p["start"] >= 0.55 * 650 - 60 for p in picks)
    assert all("自媒体" in p["why"] for p in picks)
    mid = R.focus_ranges(sents, 650, R.focus_of("中间的一些感悟"), max_s=90)
    assert mid and 200 <= mid[0]["start"] <= 400


# --------------------------------------------------------------------------- revise
def test_revise_rules(tmp_path):
    a = build("mixed", tmp_path)
    plan = PL.make_plan("这些素材都帮我做成内容，发小红书和抖音", analysis=a, asr="off")
    assert set(plan["projects"][0]["params"]["platforms"]) == {"xiaohongshu:vertical", "douyin"}
    plan["analysis"]["inputs"] = [str(tmp_path / "course")]
    # revise re-reads the analysis from the inputs: give it the same fake one
    orig = I.analyze
    I.analyze = lambda *a_, **k: a
    try:
        p2 = PL.revise(plan, "只要小红书")
        assert all(all(x.startswith("xiaohongshu") for x in p["params"]["platforms"]) for p in p2["projects"])
        p3 = PL.revise(p2, "每条 60 秒内")
        lf = next(p for p in p3["projects"] if p["recipe"] == "longform-to-short")
        assert lf["params"]["max_s"] == 60 and lf["params"]["min_s"] <= 60
        p4 = PL.revise(p3, "不要讲解视频")
        assert "explainer" not in [p["recipe"] for p in p4["projects"]]
        assert [r["prompt"] for r in p4["revisions"]] == ["只要小红书", "每条 60 秒内", "不要讲解视频"]
        assert PL.validate(p4) == []
        p5 = PL.revise(p4, "嗯嗯")
        assert any("没看懂" in w for w in p5["warnings"]) and len(p5["projects"]) == len(p4["projects"])
    finally:
        I.analyze = orig


def test_plan_stores_platforms_international_first(tmp_path):
    a = build("talk", tmp_path)
    plan = PL.make_plan("这段口播剪干净，发抖音、小红书和 TikTok", analysis=a, asr="off")
    plats = plan["projects"][0]["params"]["platforms"]
    assert [x.split(":")[0] for x in plats] == ["tiktok", "xiaohongshu", "douyin"]

    def call(system, prompt):        # a model answer in Chinese-first order is stored in registry order too
        cur = plan["projects"][0]
        return dict(json=dict(projects=[dict(recipe="talkinghead", name=cur["name"], inputs=cur["inputs"],
                                             items=dict(method="per-file"),
                                             params=dict(cur["params"], platforms=["douyin", "youtube-shorts"]))],
                              summary_zh="ok"), model="mock")
    p2 = PL.make_plan("这段口播剪干净", analysis=a, asr="off", call=call)
    assert [x.split(":")[0] for x in p2["projects"][0]["params"]["platforms"]] == ["youtube-shorts", "douyin"]
    assert [x.split(":")[0] for x in p2["projects"][0]["outputs"]["platforms"]] == ["youtube-shorts", "douyin"]


def test_revise_with_model(tmp_path):
    a = build("talk", tmp_path)
    plan = PL.make_plan("这段口播剪干净", analysis=a, asr="off")
    orig = I.analyze
    I.analyze = lambda *a_, **k: a

    def call(system, body):
        assert "FOLLOW-UP INSTRUCTION: 也发抖音" in body and "CURRENT PLAN" in body
        cur = plan["projects"][0]
        return dict(json=dict(projects=[dict(recipe="talkinghead", name=cur["name"], inputs=cur["inputs"],
                                             items=dict(method="per-file"),
                                             params=dict(cur["params"], platforms=["xiaohongshu:full", "douyin"]))],
                              summary_zh="加上抖音"), model="mock")
    try:
        p2 = PL.revise(plan, "也发抖音", call=call)
    finally:
        I.analyze = orig
    # the model's shape guess (":full") is not binding: the persona's 小红书 shape (3:4) stays
    assert p2["projects"][0]["params"]["platforms"] == ["xiaohongshu:vertical", "douyin"]
    assert p2["summary_zh"] == "加上抖音" and p2["planner"]["fallback"] is False


# --------------------------------------------------------------------------- apply
@pytest.fixture(scope="module")
def tiny(tmp_path_factory):
    if not HAS_FFMPEG:
        pytest.skip("ffmpeg not installed")
    import _batch_helpers as H
    x, truth, dur = H.synth_speech([("大家", .4), ("好", .3), ("今天", .4)])
    return H.make_video(str(tmp_path_factory.mktemp("tiny") / "t.mp4"), x, 48000, dur)


def test_apply_mixed_creates_series_and_projects(tmp_path, tiny):
    a = build("mixed", tmp_path / "m", tiny)
    plan = PL.make_plan("这些素材都帮我做成内容", analysis=a, asr="off")
    dry = AP.apply_plan(plan, str(tmp_path / "out"), dry_run=True)
    assert dry["dry_run"] and not os.path.exists(tmp_path / "out")
    res = AP.apply_plan(plan, str(tmp_path / "out"))
    assert res["ok"] and res["series"] and len(res["projects"]) == 3
    from vstudio.project import home as HM
    s = HM.series_view(res["series"])
    assert len(s["projects"]) == 3
    lf = Project(res["projects"][0]["dir"])
    assert lf.data["recipe"] == "longform-to-short" and lf.data["items"] == [] and lf.data["series"] == res["series"]
    ex = Project(res["projects"][1]["dir"])
    assert ex.data["recipe"] == "explainer" and len(ex.data["items"]) == 3
    assert os.path.exists(os.path.join(res["dir"], "intake.json"))
    assert all("--pilot" in c["run"] for c in res["projects"])


def test_apply_focus_rows_become_ranged_items(tmp_path, tiny):
    a = build("finished", tmp_path / "m", tiny)
    plan = PL.make_plan("把后面自媒体的思考单独剪出来", analysis=a, asr="off")
    res = AP.apply_plan(plan, str(tmp_path / "out"))
    pr = Project(res["projects"][0]["dir"])
    assert pr.data["recipe"] == "talkinghead" and pr.data["params"]["layout"] == "band"
    assert pr.data["params"]["captions"] is True
    assert all(i["params"]["range"][0] >= 450 and i["inputs"]["video"].endswith("_final.mp4") for i in pr.data["items"])
    assert res["series"] is None


def test_apply_rejects_invalid_plan(tmp_path):  # noqa: D103
    a = build("talk", tmp_path)
    plan = PL.make_plan("这段口播剪干净", analysis=a, asr="off")
    plan["projects"][0]["recipe"] = "invented"
    with pytest.raises(AP.ApplyError):
        AP.apply_plan(plan, str(tmp_path / "out"))


# --------------------------------------------------------------------------- inventory on real (synthetic) files
def test_docs_extract(tmp_path):
    md = tmp_path / "notes.md"
    md.write_text("# 三个误区\n- 第一 不要追热点\n- 第二 先写稿\n- 第三 每周复盘\n- 第四 看数据\n参考 https://example.com/a\n",
                  encoding="utf-8")
    f = D.analyze(str(md))
    assert f["title"] == "三个误区" and f["urls"] == ["https://example.com/a"] and f["lang"] == "zh"
    assert f["shape"] in ("notes", "outline")
    srt = tmp_path / "a.srt"
    srt.write_text("1\n00:00:01,000 --> 00:00:03,500\n大家好\n\n2\n00:01:00,000 --> 00:01:02,000\n再见\n", encoding="utf-8")
    s = D.analyze(str(srt))
    assert s["shape"] == "subtitles" and s["duration"] == 62.0
    script = tmp_path / "s.txt"
    script.write_text("第1集\n第一场 海边 日\n小林：你来了。\n阿月：嗯。\n第二场 屋内 夜\n小林：明天走吗？\n阿月：不走了。\n"
                      "第三场 车站\n小林：保重。\n第2集\n阿月：我回来了。\n", encoding="utf-8")
    assert D.analyze(str(script))["shape"] == "script"
    try:
        import docx
    except ImportError:
        docx = None
    if docx:
        d = docx.Document()
        d.add_heading("课程讲义", 1)
        d.add_paragraph("第一部分：选题")
        p = tmp_path / "x.docx"
        d.save(str(p))
        assert D.analyze(str(p))["title"] == "课程讲义"
    pptx_missing = tmp_path / "deck.pptx"
    pptx_missing.write_bytes(b"not a real deck")
    r = D.analyze(str(pptx_missing))
    assert r["chars"] is None and r["note"]


@pytest.mark.skipif(not HAS_FFMPEG, reason="ffmpeg not installed")
def test_inventory_synthetic(tmp_path, monkeypatch):
    import _batch_helpers as H
    from PIL import Image
    d = tmp_path / "in"
    (d / "photos").mkdir(parents=True)
    x, truth, dur = H.synth_speech([("大家", .4), ("好", .3), ("今天", .4), ("讲", .3)])
    H.make_video(str(d / "talk.mp4"), x, 48000, dur)
    for k in range(3):
        Image.new("RGB", (64, 48), (k * 40, 90, 70)).save(d / "photos" / f"p{k}.jpg")
    (d / "brief.md").write_text("# 选题\n看 https://example.com\n", encoding="utf-8")
    (d / ".hidden.mp4").write_bytes(b"x")
    (d / "x.bin").write_bytes(b"x")
    calls = []
    monkeypatch.setattr(I, "TRANSCRIBE", lambda wav, lang: calls.append(wav) or dict(
        language="zh", segments=[dict(start=0, end=1, text="大家好今天讲一个方法")], words=[{}] * 6))
    monkeypatch.setattr(I, "FACES", lambda img: [0.05])
    a = I.analyze([str(d)])
    by = {f["rel"]: f for f in a["files"]}
    assert ".hidden.mp4" not in by and by["x.bin"]["kind"] == "other"
    v = by["talk.mp4"]
    assert v["speech"] is True and v["orientation"] == "vertical" and v["talking_head"] is True and v["qhash"]
    assert I.material_role(v) == "talking-head"
    assert by["brief.md"]["urls"] == ["https://example.com"] and a["urls"] == ["https://example.com"]
    assert a["groups"] and a["groups"][0]["kind"] == "photo-set" and a["groups"][0]["n"] == 3
    n = len(calls)
    a2 = I.analyze([str(d)])
    assert len(calls) == n and all(f.get("cached") for f in a2["files"] if f["kind"] != "other")
    full = I.analyze([str(d / "talk.mp4")], asr="full")
    tr = I.load_transcript(full["files"][0])
    assert tr["sentences"][0]["text"].startswith("大家好")
    c = I.compact(a)
    assert all("path" not in f for f in c["files"])
    assert not any(n.endswith(".asr.json") for n in os.listdir(d))       # inputs stay untouched


@pytest.mark.skipif(not HAS_FFMPEG, reason="ffmpeg not installed")
def test_inventory_cache_keeps_languages_apart(tmp_path, monkeypatch):
    """An English pass of a file analyzed before (auto-detect / Chinese) transcribes again, then is cached itself."""
    import _batch_helpers as H
    monkeypatch.setenv("VSTUDIO_CACHE", str(tmp_path / "cache"))
    x, _truth, dur = H.synth_speech([("hello", .4), ("everyone", .5), ("today", .4)])
    H.make_video(str(tmp_path / "talk.mp4"), x, 48000, dur)
    calls = []
    monkeypatch.setattr(I, "TRANSCRIBE", lambda wav, lang: calls.append(lang) or dict(
        language=lang or "zh", segments=[dict(start=0, end=1, text="hello everyone today")], words=[{}] * 3))
    monkeypatch.setattr(I, "FACES", lambda img: [0.05])
    src = [str(tmp_path / "talk.mp4")]
    assert I.analyze(src)["files"][0]["language"] == "zh"
    n = len(calls)
    en = I.analyze(src, language="en")["files"][0]
    assert len(calls) > n and en["language"] == "en" and not en["cached"]
    n = len(calls)
    assert I.analyze(src, language="en")["files"][0]["cached"] and len(calls) == n
    assert I.analyze(src)["files"][0]["cached"] and I.analyze(src)["files"][0]["language"] == "zh"


def test_visual_facts_heuristics():
    th = [dict(faces=1, face_area=0.14, edges=0.07, sat=70, flat=0.5, cap_low=40, cap_mid=0.5, cap_top=0.4)] * 6
    v = I.visual_facts(th, 1920, 1080)
    assert v["talking_head"] and v["burned_captions"] and not v["screen_share"] and not v["multi_person"]
    scr = [dict(faces=0, face_area=None, edges=0.04, sat=3, flat=0.8, cap_low=10, cap_mid=2, cap_top=10)] * 6
    v = I.visual_facts(scr, 1280, 720)
    assert v["screen_share"] and not v["burned_captions"] and not v["talking_head"]
    call = [dict(faces=3, face_area=0.02, edges=0.05, sat=60, flat=0.5, cap_low=0, cap_mid=0, cap_top=0)] * 6
    assert I.visual_facts(call, 1280, 720)["multi_person"]


# --------------------------------------------------------------------------- CLI
def _cli(*args):
    e = dict(os.environ)
    e["PYTHONPATH"] = os.path.join(ROOT, "lib") + os.pathsep + e.get("PYTHONPATH", "")
    return subprocess.run([sys.executable, "-m", "vstudio.intake", *args], capture_output=True, text=True, env=e)


def test_cli_plan_revise_apply(tmp_path):
    a = build("pdf", tmp_path / "m")
    ap = tmp_path / "analysis.json"
    ap.write_text(json.dumps(a, ensure_ascii=False), encoding="utf-8")
    r = _cli("schema", "--json")
    assert r.returncode == 0 and json.loads(r.stdout)["$id"] == "vstudio.intake.plan"
    pj = tmp_path / "plan.json"
    r = _cli("plan", "--prompt", "用这份 PDF 做 5 条讲解短视频", "--analysis", str(ap), "--provider", "none", "--asr", "off",
             "--json", "--quiet", "--out", str(pj))
    assert r.returncode == 0, r.stderr
    plan = json.loads(r.stdout)
    assert plan["projects"][0]["recipe"] == "explainer" and json.loads(pj.read_text())["id"] == plan["id"]
    r = _cli("apply", "--plan", str(pj), "--out", str(tmp_path / "out"), "--json")
    assert r.returncode == 0, r.stderr
    res = json.loads(r.stdout)
    assert res["ok"] and len(res["projects"][0]["items"]) == 5
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps(dict(plan, projects=[])), encoding="utf-8")
    r = _cli("apply", "--plan", str(bad), "--json")
    assert r.returncode == 5 and json.loads(r.stdout)["ok"] is False


def test_intake_cli_attempt_is_90s(monkeypatch):
    """Each CLI provider attempt gets 90 s (the same per-attempt policy as Create), then the chain's fallback."""
    seen = {}

    def complete(task, system, body, **kw):
        seen.update(kw)
        raise PL.LLM.LLMError("stop here")
    monkeypatch.delenv("VSTUDIO_LLM_CLI_TIMEOUT", raising=False)
    monkeypatch.setattr(PL.LLM, "complete", complete)
    monkeypatch.setattr(PL.LLM, "route", lambda *a, **k: type("R", (), dict(provider="claude-code", model=None,
                                                                             source="test"))())
    monkeypatch.setattr(PL, "_prompt_doc", lambda *a, **k: "request")
    js, info = PL._call_model("剪干净", {}, {}, {})
    assert js is None and info["fallback"] and seen["cli_timeout"] == 90


def test_intake_cli_attempt_grows_with_the_prompt(monkeypatch):
    """A 1,000-file intake with a 12-minute transcript (~95k chars) took Claude Code ~140 s: 90 s cut it off and Codex
    answered. The per-attempt limit grows with the prompt (capped)."""
    seen = {}

    def complete(task, system, body, **kw):
        seen.update(kw)
        raise PL.LLM.LLMError("stop here")
    monkeypatch.delenv("VSTUDIO_LLM_CLI_TIMEOUT", raising=False)
    monkeypatch.setattr(PL.LLM, "complete", complete)
    monkeypatch.setattr(PL.LLM, "route", lambda *a, **k: type("R", (), dict(provider="claude-code", model=None,
                                                                             source="test"))())
    monkeypatch.setattr(PL, "_prompt_doc", lambda *a, **k: "x" * 90000)
    js, info = PL._call_model("剪干净", {}, {}, {})
    assert seen["cli_timeout"] >= 240 and info["cli_timeout"] == seen["cli_timeout"]
    assert PL.cli_timeout_for(10 ** 7) == PL.MAX_CLI_TIMEOUT and PL.cli_timeout_for(5000) == 90
